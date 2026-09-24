"""KPIs e breakdowns do módulo de RH — usado pelo dashboard e pelo relatório impresso."""
import logging
from collections import Counter, defaultdict
from typing import Optional

from fastapi import APIRouter, Depends, Query

from auth import require_role
from db import get_supabase

router = APIRouter(prefix="/api/rh/dashboard")
log = logging.getLogger(__name__)

_ROLES = ("admin", "rh")


def _require_rh(user=Depends(require_role(*_ROLES))):
    return user


@router.post("/relatorio-semanal/enviar-agora")
def enviar_relatorio_semanal_agora(user=Depends(_require_rh)):
    from services.relatorio_semanal import gerar_e_enviar
    ok = gerar_e_enviar()
    return {"ok": ok}




def _linhas_filtradas(sb, filtros: dict) -> list[dict]:
    from routes.vagas import _FILTER_COLS, _SELECT, _serialize

    from services.paginacao import buscar_todos

    def _query():
        query = sb.table("rh_vagas").select(_SELECT)
        if filtros.get("status_id"):
            query = query.in_("status_id", filtros["status_id"])
        if filtros.get("data_inicio"):
            query = query.gte("data_recebimento", filtros["data_inicio"])
        if filtros.get("data_fim"):
            query = query.lte("data_recebimento", filtros["data_fim"])
        for col in _FILTER_COLS:
            if filtros.get(col):
                query = query.eq(col, filtros[col])
        return query

    rows = [_serialize(r) for r in buscar_todos(_query)]

    if filtros.get("ano"):
        anos_str = {str(a) for a in filtros["ano"]}
        rows = [r for r in rows if (r.get("data_recebimento") or "")[:4] in anos_str]

    q = filtros.get("q")
    if q:
        q_lower = q.lower()
        rows = [
            r for r in rows
            if q_lower in (r.get("candidato") or "").lower()
            or q_lower in (r.get("cargo") or "").lower()
            or q_lower in (r.get("numero_requisicao") or "").lower()
            or q_lower in (r.get("responsavel") or "").lower()
        ]
    return rows


def _filtros(
    q: Optional[str] = Query(None),
    status_id: Optional[list[str]] = Query(None),
    ano: Optional[list[int]] = Query(None),
    data_inicio: Optional[str] = Query(None),
    data_fim: Optional[str] = Query(None),
    empresa_id: Optional[str] = Query(None),
    tipo_vaga_id: Optional[str] = Query(None),
    tipo_contrato_id: Optional[str] = Query(None),
    nivel_id: Optional[str] = Query(None),
    hierarquia_id: Optional[str] = Query(None),
    etapa_atual_id: Optional[str] = Query(None),
    secao_id: Optional[str] = Query(None),
    responsavel_id: Optional[str] = Query(None),
    requisitante_id: Optional[str] = Query(None),
    cargo_id: Optional[str] = Query(None),
) -> dict:
    return dict(locals())


def _pct(parte: int, total: int) -> Optional[float]:
    return round(100 * parte / total, 1) if total else None


def _media(valores: list) -> Optional[float]:
    valores = [v for v in valores if v is not None]
    return round(sum(valores) / len(valores), 1) if valores else None


def _resumo_fase(rows: list[dict], fase: str) -> dict:
    from services.sla import AVALIAVEIS, NO_PRAZO_SET

    por_status = Counter(r["sla"][fase]["status"] for r in rows)
    avaliaveis = [r for r in rows if r["sla"][fase]["status"] in AVALIAVEIS]
    no_prazo = [r for r in avaliaveis if r["sla"][fase]["status"] in NO_PRAZO_SET]
    concluidas = [r for r in avaliaveis if r["sla"][fase]["fim"]]
    return {
        "por_status": dict(por_status),
        "avaliadas": len(avaliaveis),
        "pct_no_prazo": _pct(len(no_prazo), len(avaliaveis)),
        "em_andamento_no_prazo": por_status.get("NO PRAZO", 0),
        "em_andamento_atrasadas": por_status.get("ATRASADO", 0),
        "concluidas_no_prazo": por_status.get("CONCLUÍDA NO PRAZO", 0),
        "concluidas_com_atraso": por_status.get("CONCLUÍDA COM ATRASO", 0),
        "media_dias_concluidas": _media([r["sla"][fase]["dias"] for r in concluidas]),
        "media_sla": _media([r["sla"][fase]["sla"] for r in avaliaveis]),
        "estimadas": sum(1 for r in avaliaveis if r["sla"][fase].get("estimado")),
    }


def _linha_relatorio(r: dict) -> dict:
    s = r["sla"]
    return {
        "id": r.get("id"), "numero_requisicao": r.get("numero_requisicao"),
        "cargo": r.get("cargo"), "empresa": r.get("empresa"), "nivel": r.get("nivel"),
        "responsavel": r.get("responsavel"), "requisitante": r.get("requisitante"),
        "status": r.get("status"), "etapa_atual": r.get("etapa_atual"),
        "rs": s["rs"], "adm": s["adm"], "etapa": s["etapa"],
    }


def _bloco_sla(sb, rows: list[dict]) -> dict:
    from services.sla import ATRASADO, AVALIAVEIS, NO_PRAZO_SET, referencias

    _, etapas = referencias(sb)

    # % no prazo por recrutador / nível
    def _agrupar(chave: str) -> list[dict]:
        grupos: dict[str, list] = defaultdict(list)
        for r in rows:
            grupos[r.get(chave) or "NÃO INFORMADO"].append(r)
        saida = []
        for nome, itens in grupos.items():
            item = {"nome": nome, "total": len(itens)}
            for fase in ("rs", "adm"):
                av = [i for i in itens if i["sla"][fase]["status"] in AVALIAVEIS]
                ok = [i for i in av if i["sla"][fase]["status"] in NO_PRAZO_SET]
                item[f"{fase}_avaliadas"] = len(av)
                item[f"{fase}_pct_no_prazo"] = _pct(len(ok), len(av))
                item[f"{fase}_atrasadas"] = sum(1 for i in itens if i["sla"][fase]["status"] == ATRASADO)
            saida.append(item)
        return sorted(saida, key=lambda x: -x["total"])

    # Prazo por etapa: vagas paradas agora em cada etapa ativa + média histórica
    ativas = sorted(
        [e for e in etapas.values() if e.get("ativo") and e.get("fase") in ("RS", "ADMISSAO")],
        key=lambda e: e["ordem"],
    )
    ids_rows = {r["id"] for r in rows}
    from services.paginacao import buscar_todos

    hist = buscar_todos(lambda: sb.table("rh_vagas_etapas_hist").select("id,vaga_id,etapa_id,inicio,fim"))
    from datetime import date as _date
    hoje = _date.today()
    dur_por_etapa: dict[str, list[int]] = defaultdict(list)
    for h in hist:
        if h["vaga_id"] not in ids_rows or not h.get("fim"):
            continue
        ini = _date.fromisoformat(h["inicio"])
        fim = _date.fromisoformat(h["fim"])
        dur_por_etapa[h["etapa_id"]].append(max((fim - ini).days, 0))

    em_aberto = [r for r in rows if r.get("status_em_aberto")]
    etapas_saida = []
    for e in ativas:
        atuais = [r for r in em_aberto if r.get("etapa_atual_id") == e["id"]]
        etapas_saida.append({
            "etapa": e["nome"], "ordem": e["ordem"], "fase": e["fase"], "externa": bool(e.get("externa")),
            "sla": 3 if e.get("externa") else None,
            "qtd_atual": len(atuais),
            "dias_medio_atual": _media([r["sla"]["etapa"]["dias"] for r in atuais]),
            "estouradas": sum(1 for r in atuais if r["sla"]["etapa"]["status"] == ATRASADO),
            "sem_data_inicio": sum(1 for r in atuais if r["sla"]["etapa"]["status"] == "INFORMAR DATA DE INÍCIO"),
            "dias_medio_historico": _media(dur_por_etapa.get(e["id"], [])),
            "amostras_historico": len(dur_por_etapa.get(e["id"], [])),
        })

    atrasadas_rs = [r for r in rows if r["sla"]["rs"]["status"] == ATRASADO]
    atrasadas_adm = [r for r in rows if r["sla"]["adm"]["status"] == ATRASADO]
    etapa_estourada = [r for r in rows if r["sla"]["etapa"]["status"] == ATRASADO]
    return {
        "rs": _resumo_fase(rows, "rs"),
        "adm": _resumo_fase(rows, "adm"),
        "por_recrutador": _agrupar("responsavel"),
        "por_nivel": _agrupar("nivel"),
        "etapas": etapas_saida,
        "atrasadas_rs": [_linha_relatorio(r) for r in atrasadas_rs],
        "atrasadas_adm": [_linha_relatorio(r) for r in atrasadas_adm],
        "etapa_estourada": [_linha_relatorio(r) for r in etapa_estourada],
    }


@router.get("/sla-relatorio")
def sla_relatorio(
    formato: str = Query("json", pattern="^(json|xlsx)$"),
    filtros: dict = Depends(_filtros),
    user=Depends(_require_rh),
):
    """Relatório vaga a vaga de SLA por fase (R&S / Admissão / etapa atual)."""
    rows = _linhas_filtradas(get_supabase(), filtros)
    linhas = [_linha_relatorio(r) for r in rows]
    linhas.sort(key=lambda x: (x["rs"]["inicio"] or ""), reverse=True)
    if formato == "json":
        return linhas

    import io

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "SLA por vaga"
    cab = [
        "Nº REQUISIÇÃO", "CARGO", "EMPRESA", "NÍVEL", "RECRUTADOR", "STATUS", "ETAPA ATUAL",
        "R&S INÍCIO", "R&S SLA (DIAS)", "R&S LIMITE", "R&S FIM", "R&S DIAS", "R&S STATUS",
        "ADM INÍCIO", "ADM SLA (DIAS)", "ADM LIMITE", "ADM FIM", "ADM DIAS", "ADM STATUS",
        "ETAPA DIAS", "ETAPA SLA", "ETAPA STATUS",
    ]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="00694E")
    for l in linhas:
        rs, adm, et = l["rs"], l["adm"], l["etapa"]
        ws.append([
            l["numero_requisicao"], l["cargo"], l["empresa"], l["nivel"], l["responsavel"], l["status"], l["etapa_atual"],
            rs["inicio"], rs["sla"], rs["limite"], rs["fim"], rs["dias"], rs["status"] + (" (estimado)" if rs.get("estimado") else ""),
            adm["inicio"], adm["sla"], adm["limite"], adm["fim"], adm["dias"], adm["status"],
            et["dias"], et["sla"], et["status"],
        ])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, min(40, max(len(str(c.value or "")) for c in col) + 2))
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="relatorio_sla_vagas.xlsx"'},
    )


@router.get("")
def dashboard(filtros: dict = Depends(_filtros), user=Depends(_require_rh)):
    sb = get_supabase()
    rows = _linhas_filtradas(sb, filtros)

    abertas = [r for r in rows if r.get("status_em_aberto")]
    concluidas = [r for r in rows if r.get("status_concluido")]
    canceladas = [r for r in rows if r.get("status") == "CANCELADO"]
    congeladas = [r for r in rows if r.get("status") == "CONGELADO"]

    slas_validos = [r["sla_ok"] for r in rows if r.get("sla_ok") is not None]
    pct_no_prazo = round(100 * sum(slas_validos) / len(slas_validos), 1) if slas_validos else None

    dias_validos = [r["dias_corridos"] for r in concluidas if r.get("dias_corridos") is not None]
    sla_medio_dias = round(sum(dias_validos) / len(dias_validos), 1) if dias_validos else None

    atrasadas = [r for r in abertas if r.get("sla_ok") is False]

    # SLA estourado (já passou do prazo) e estourando (faltam <=3 dias, ainda no prazo)
    sla_estourado = [r for r in atrasadas]
    sla_estourando = []
    for r in abertas:
        dias = r.get("dias_corridos")
        alvo = r["sla"]["rs"]["sla"]
        if dias is None or not alvo or r.get("sla_ok") is False:
            continue
        if alvo - dias <= 3:
            sla_estourando.append(r)

    def _resumo_alerta(r):
        return {
            "id": r.get("id"), "numero_requisicao": r.get("numero_requisicao"),
            "cargo": r.get("cargo"), "empresa": r.get("empresa"),
            "responsavel": r.get("responsavel"), "dias_corridos": r.get("dias_corridos"),
            "sla_alvo_dias": r["sla"]["rs"]["sla"], "etapa_atual": r.get("etapa_atual"),
            "status": r.get("status"),
        }

    # Vagas por analista
    por_analista_map: dict[str, dict] = {}
    for r in rows:
        nome = r.get("responsavel") or "NÃO INFORMADO"
        item = por_analista_map.setdefault(nome, {
            "analista": nome, "total": 0, "abertas": 0, "concluidas": 0,
            "canceladas": 0, "congeladas": 0,
        })
        item["total"] += 1
        if r.get("status_em_aberto"):
            item["abertas"] += 1
        if r.get("status_concluido"):
            item["concluidas"] += 1
        if r.get("status") == "CANCELADO":
            item["canceladas"] += 1
        if r.get("status") == "CONGELADO":
            item["congeladas"] += 1

    por_status = Counter(r.get("status") or "NÃO INFORMADO" for r in rows)
    por_empresa = Counter(r.get("empresa") or "NÃO INFORMADO" for r in rows)
    por_empresa_fechadas = Counter(r.get("empresa") or "NÃO INFORMADO" for r in concluidas)
    top_cargos = Counter(r.get("cargo") for r in rows if r.get("cargo"))

    tendencia: dict[str, dict[str, int]] = defaultdict(lambda: {"abertas": 0, "concluidas": 0})
    for r in rows:
        mes = (r.get("data_recebimento") or "")[:7]
        if mes:
            tendencia[mes]["abertas"] += 1
    for r in concluidas:
        mes = (r.get("data_admissao") or r.get("data_recebimento") or "")[:7]
        if mes:
            tendencia[mes]["concluidas"] += 1

    etapas_resp = sb.table("rh_etapas_processo").select("id,nome,ordem").eq("ativo", True).order("ordem").execute()
    contagem_etapa = Counter(r.get("etapa_atual_id") for r in rows if r.get("etapa_atual_id"))
    funil_etapas = [
        {"etapa": e["nome"], "ordem": e["ordem"], "total": contagem_etapa.get(e["id"], 0)}
        for e in (etapas_resp.data or [])
    ]

    return {
        "sla_fases": _bloco_sla(sb, rows),
        "kpis": {
            "total": len(rows),
            "abertas": len(abertas),
            "concluidas_periodo": len(concluidas),
            "sla_medio_dias": sla_medio_dias,
            "pct_no_prazo": pct_no_prazo,
            "atrasadas": len(atrasadas),
            "canceladas": len(canceladas),
            "congeladas": len(congeladas),
        },
        "por_status": [{"status": k, "total": v} for k, v in sorted(por_status.items(), key=lambda x: -x[1])],
        "por_empresa": [{"empresa": k, "total": v} for k, v in sorted(por_empresa.items(), key=lambda x: -x[1])],
        "por_empresa_fechadas": [{"empresa": k, "total": v} for k, v in sorted(por_empresa_fechadas.items(), key=lambda x: -x[1])],
        "top_cargos": [{"cargo": k, "total": v} for k, v in top_cargos.most_common(10)],
        "tendencia_mensal": [
            {"mes": mes, **vals} for mes, vals in sorted(tendencia.items())
        ],
        "funil_etapas": funil_etapas,
        "por_analista": sorted(por_analista_map.values(), key=lambda x: -x["total"]),
        "sla_estourado": [_resumo_alerta(r) for r in sla_estourado],
        "sla_estourando": [_resumo_alerta(r) for r in sla_estourando],
        "abertas_lista": [_resumo_alerta(r) for r in abertas],
        "canceladas_congeladas_lista": [_resumo_alerta(r) for r in canceladas + congeladas],
    }
